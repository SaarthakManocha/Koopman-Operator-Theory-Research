"""post_docx.py - page margins, centred figures, ruled tables with a shaded header row."""
import sys
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
f = sys.argv[1]; d = Document(f)
for s in d.sections:
    s.left_margin = s.right_margin = Inches(0.9); s.top_margin = s.bottom_margin = Inches(0.9)
maxw = Inches(6.5)
for sh in d.inline_shapes:
    if sh.width > maxw:
        r = maxw / sh.width; sh.width = int(sh.width * r); sh.height = int(sh.height * r)
for p in d.paragraphs:
    if p._p.xpath('.//pic:pic') or p.style.name in ("Image Caption", "Captioned Figure", "Title", "Author", "Date"):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if p.text.startswith("TABLE "):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs: r.font.size = Pt(9.5)
import re
inrefs = False
for p in d.paragraphs:
    if p.text.strip() == "References": inrefs = True; continue
    if inrefs and re.match(r"^\d+\]", p.text) and p.runs:
        p.runs[0].text = "[" + p.runs[0].text
for t in d.tables:
    tblPr = t._tbl.tblPr
    for old in tblPr.findall(qn('w:tblBorders')): tblPr.remove(old)
    b = OxmlElement('w:tblBorders')
    for e in ['top', 'bottom', 'insideH']:
        x = OxmlElement('w:' + e); x.set(qn('w:val'), 'single'); x.set(qn('w:sz'), '4'); x.set(qn('w:color'), '888888'); b.append(x)
    tblPr.append(b)
    jc = OxmlElement('w:jc'); jc.set(qn('w:val'), 'center'); tblPr.append(jc)
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs: r.font.size = Pt(9)
    for c in t.rows[0].cells:
        for p in c.paragraphs:
            for r in p.runs: r.bold = True
        sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:fill'), 'EEF2F7'); c._tc.get_or_add_tcPr().append(sh)
d.save(f); print("figures", len(d.inline_shapes), "tables", len(d.tables))
